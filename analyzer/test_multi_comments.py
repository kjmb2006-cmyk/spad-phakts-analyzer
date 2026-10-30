#!/usr/bin/env python3
"""
Test script — lecture et interprétation de la comparaison multi-enquête.

Sous les données de chaque comparaison, la page « Analyse multi-enquête » affiche une
« Lecture automatique » (ce que l'on observe, si l'écart est établi, de combien, avec
quelles précautions) et, au-dessus des comparaisons, une synthèse. Vérifie :
  - le texte dit ce que les chiffres disent (écart établi, stable, à confirmer) ;
  - il nuance (effectifs faibles, tableau creux, distribution asymétrique) ;
  - la page l'affiche dans les deux modes (indicateurs détectés / panier manuel) et
    conserve le commentaire libre de l'utilisateur.
"""
import io
import os
import shutil
import subprocess
import tempfile

import numpy as np
import pandas as pd

import app as m
from modules import multi_comments as mc
from modules.multi_survey import apply_fdr, compare_categorical_by_survey, cross_survey_test, evolution_categorical, merge_surveys

print("=" * 70)
print("TEST — lecture et interprétation de la comparaison multi-enquête")
print("=" * 70)

checks = 0


def ok(cond, msg=''):
    global checks
    assert cond, msg
    checks += 1


rng = np.random.default_rng(7)
N = 200


def survey(oui, age, sexe_f=0.5, n=N):
    return pd.DataFrame({
        'Fumeur__B': np.where(rng.random(n) < oui, 'OUI', 'NON'),
        'Sexe__X': np.where(rng.random(n) < sexe_f, 'Femme', 'Homme'),
        'Age__1Y': rng.normal(age, 6, n).round(0),
    })


surveys = {'Enquête 2023': survey(0.20, 35), 'Enquête 2024': survey(0.40, 40), 'Enquête 2025': survey(0.60, 45)}
order = list(surveys)
merged = merge_surveys(surveys, mode='phakts')


def compare(variable, kind):
    r = (compare_categorical_by_survey(merged, variable, order) if kind == 'categorical'
         else m.compare_continuous_by_survey(merged, variable, order))
    r['kind'] = kind
    r['variable'] = variable
    r['evolution'] = (evolution_categorical if kind == 'categorical' else m.evolution_continuous)(merged, variable, order)
    r['test'] = cross_survey_test(merged, variable, kind)
    return r


results = [compare('Fumeur__B', 'categorical'), compare('Sexe__X', 'categorical'), compare('Age__1Y', 'continuous')]
apply_fdr(results)
fumeur, sexe, age = results

# ── Le test d'écart porte l'intensité et la fiabilité du χ²
ok(0 < fumeur['test']['cramers_v'] <= 1 and fumeur['test']['low_expected_pct'] == 0.0, fumeur['test'])
ok(fumeur['stats']['surveys'][0] == {'name': 'Enquête 2023', 'n': N, 'top': 'NON', 'top_pct': fumeur['stats']['surveys'][0]['top_pct']})
ok(age['stats']['surveys'][2]['mean'] > age['stats']['surveys'][0]['mean'])
print("OK — stats par enquête, V de Cramér et cases à effectif théorique faible")

# ── Écart établi : la lecture le dit, avec son intensité et les évolutions
txt = mc.interpret(fumeur)
ok('diffèrent d\'une enquête à l\'autre' in txt and 'V de Cramér' in txt, txt)
ok('passe de' in txt and 'IC à 95 %' in txt and '+' in txt, txt)
ok('Précaution' not in txt, txt)                   # effectifs suffisants : rien à nuancer
print("OK — écart établi : décrit, chiffré, avec amplitude et intervalle de confiance")

# ── Indicateur stable : on ne conclut pas à l'équivalence
txt = mc.interpret(sexe)
ok('aucune différence' in txt.lower() or 'significatif avant correction' in txt, txt)
ok('marge du hasard' in txt and 'change d\'une enquête' not in txt, txt)   # 52 % / 52 % : pas de « changement » à commenter
untested = dict(sexe, test=None)
ok('marge du hasard' not in mc.interpret(untested) and 'non réalisée' in mc.interpret(untested), mc.interpret(untested))
print("OK — indicateur stable : pas d'écart établi")

# ── Continue : moyennes, test, tendance
txt = mc.interpret(age)
ok('moyenne' in txt and 'médiane' in txt and "La moyenne augmente d'une enquête à l'autre" in txt, txt)
ok(age['test']['test'] in txt, txt)
print("OK — variable continue : moyenne, médiane, tendance")

# ── Chiffres à la française
ok('p = 0.' not in mc.interpret(sexe) and 'p < 0.' not in mc.interpret(fumeur), mc.interpret(sexe))
ok(mc._fr(3.14159, 2) == '3,14' and mc._signed(-2.5) == '−2,5' and mc._signed(0.5) == '+0,5')
print("OK — virgule décimale et signe typographique")

# ── Précautions
small = {'kind': 'categorical', 'variable': 'X',
         'stats': {'surveys': [{'name': 'A', 'n': 12, 'top': 'OUI', 'top_pct': 60.0}, {'name': 'B', 'n': 400, 'top': 'NON', 'top_pct': 70.0}],
                   'n_modalites': 20},
         'test': {'test': 'χ² (indépendance)', 'p_value': 0.01, 'q_value': 0.01, 'cramers_v': 0.4, 'low_expected_pct': 45.0},
         'evolution': []}
txt = mc.interpret(small)
ok('effectifs faibles' in txt and '« A » (n = 12)' in txt, txt)
ok('45 % des cases' in txt and 'peu fiable' in txt, txt)
ok('20 modalités' in txt, txt)
ok("la modalité dominante change" in txt, txt)
skew = {'kind': 'continuous', 'variable': 'Y', 'stats': {'surveys': [{'name': 'A', 'n': 100, 'mean': 50.0, 'median': 30.0, 'std': 20.0}]},
        'test': None, 'evolution': []}
txt = mc.interpret(skew)
ok('asymétrique' in txt and 'privilégier la médiane' in txt and 'non réalisée' in txt, txt)
ok(mc.interpret({'variable': 'Z', 'error': 'boom'}) == '')
print("OK — précautions : effectifs faibles, tableau creux, trop de modalités, asymétrie, calcul impossible")

# ── Correction des tests multiples : « significatif avant correction seulement » → à confirmer
tentative = {'kind': 'categorical', 'variable': 'T', 'stats': {'surveys': [{'name': 'A', 'n': 80, 'top': 'a', 'top_pct': 50.0}]},
             'test': {'test': 'χ² (indépendance)', 'p_value': 0.03, 'q_value': 0.2, 'cramers_v': 0.1}, 'evolution': []}
ok('à confirmer' in mc.interpret(tentative, corrected=True))
ok('à confirmer' not in mc.interpret(tentative, corrected=False) and 'diffèrent' in mc.interpret(tentative, corrected=False))
print("OK — significatif avant correction seulement : signalé « à confirmer »")

# ── Synthèse
syn = mc.synthese(results, order, corrected=True, n_errors=0)
ok('3 indicateur(s) comparé(s) entre 3 enquêtes' in syn and 'Enquête 2023 → Enquête 2024 → Enquête 2025' in syn, syn)
ok('Écarts établis' in syn and 'Fumeur__B' in syn.split('Écarts établi')[1].split('\n')[0], syn)
ok('Pour mémoire' in syn and 'populations' in syn, syn)
ok('Attention' not in syn)
ok('Attention : 2 contrôle(s)' in mc.synthese(results, order, True, n_errors=2))
ok(mc.synthese([], order) == '' and mc.synthese([{'variable': 'x', 'error': 'e'}], order) == '')
print("OK — synthèse : écarts établis, stables, avertissement sur les contrôles en erreur")

# ── Page : les deux modes
work = tempfile.mkdtemp()
FAC = ['PGT14_OS2_RC2_2__CIV_MS__1C2_Etude_Tabac__20230900__ES',
       'PGT14_OS2_RC2_2__CIV_MS__1C2_Etude_Tabac__20240900__ES',
       'PGT14_OS2_RC2_2__CIV_MS__1C2_Etude_Tabac__20250900__ES']
entries = []
for (name, df), fac in zip(surveys.items(), FAC):
    path = os.path.join(work, f'{name.replace(" ", "_")}.xlsx')
    df.to_excel(path, index=False)
    entries.append({'name': os.path.basename(path), 'path': path, 'n_obs': len(df), 'n_vars': df.shape[1],
                    'phakts_vars': df.shape[1], 'factum': fac})

m.app.config['TESTING'] = True
m.app.config['REPORTS_FOLDER'] = os.path.join(work, 'rapports')      # rien dans static/reports
os.makedirs(m.app.config['REPORTS_FOLDER'])
with m.app.test_client() as c:
    with c.session_transaction() as sess:
        sess['multi_surveys'] = entries
        sess['ms_auto_mode'] = True
        sess.pop('user_comments', None)
    page = c.get('/multi-survey').get_data(as_text=True)
    ok('Synthèse de la comparaison' in page and 'Écarts établis' in page, 'synthèse absente (mode automatique)')
    ok(page.count('Lecture automatique') >= 3, page.count('Lecture automatique'))
    ok('data-section="multi_survey"' in page and 'data-var="Fumeur__B"' in page)
    ok('V de Cramér' in page and 'IC à 95 %' in page)
    print("OK — page, mode automatique : synthèse + une lecture sous chaque comparaison")

    # le commentaire libre de l'utilisateur est repris au rechargement
    r = c.post('/api/save-comment', json={'section': 'multi_survey', 'var': 'Fumeur__B', 'text': 'Hausse liée à la campagne 2024.'})
    ok(r.status_code == 200, r.status_code)
    ok('Hausse liée à la campagne 2024.' in c.get('/multi-survey').get_data(as_text=True))
    print("OK — le commentaire libre est conservé et réaffiché")

    # mode manuel : panier de variables, une seule comparaison → pas de correction affichée
    with c.session_transaction() as sess:
        sess['ms_auto_mode'] = False
        sess['ms_compared_vars'] = ['Fumeur__B']
    page = c.get('/multi-survey').get_data(as_text=True)
    ok('Synthèse de la comparaison' in page and page.count('Lecture automatique') == 1, page.count('Lecture automatique'))
    ok('après correction des tests multiples' not in page.split('Synthèse de la comparaison')[1].split('Pour mémoire')[0], 'une seule comparaison : rien à corriger')
    with c.session_transaction() as sess:
        sess['ms_compared_vars'] = ['Fumeur__B', 'Age__1Y', 'Sexe__X']
    page = c.get('/multi-survey').get_data(as_text=True)
    ok(page.count('Lecture automatique') == 3 and 'après correction des tests multiples' in page, page.count('Lecture automatique'))
    ok(' · q ' in page and ('· p < 0,0001' in page or '· p = ' in page), 'badge de test dans l\'en-tête')
    print("OK — page, mode manuel : correction des tests multiples seulement à partir de deux comparaisons")

    # ── Rapports PDF / Word / Excel : synthèse, lecture automatique et commentaire de l'analyste
    from docx import Document
    from openpyxl import load_workbook
    ref = os.path.join(work, 'ref.xlsx')
    surveys['Enquête 2023'].to_excel(ref, index=False)
    with c.session_transaction() as sess:
        sess['data_path'] = ref
        sess['data_meta'] = {'name': 'ref', 'n_obs': N, 'n_vars': 3}
        sess['original_filename'] = 'ref.xlsx'
        sess['ms_compared_vars'] = ['Fumeur__B', 'Age__1Y']
        sess['user_comments'] = {'multi_survey:Fumeur__B': "Hausse liée à la campagne <2024> & suivi."}
    base = {'analyses': ['multi_survey'], 'report_title': 'Rapport test', 'report_author': 'Test'}

    r = c.post('/report', data=dict(base, report_format='word'))
    ok(r.status_code == 200 and r.data[:2] == b'PK', (r.status_code, r.location))
    words = '\n'.join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
    ok('Synthèse de la comparaison' in words and 'Écarts établis' in words, 'synthèse absente du Word')
    ok(words.count('Lecture automatique') == 2 and 'V de Cramér' in words, words.count('Lecture automatique'))
    ok("Commentaire de l'analyste : Hausse liée à la campagne <2024> & suivi." in words, 'commentaire libre absent du Word')

    r = c.post('/report', data=dict(base, report_format='excel'))
    ok(r.status_code == 200 and r.data[:2] == b'PK', (r.status_code, r.location))
    ws = load_workbook(io.BytesIO(r.data))['Multi-enquête']
    cells = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
    ok('Synthèse de la comparaison' in cells and cells.count('Lecture automatique :') == 2, cells[:8])
    ok(any('Hausse liée à la campagne <2024> & suivi.' in v for v in cells))

    r = c.post('/report', data=dict(base, report_format='pdf'))
    ok(r.status_code == 200 and r.data[:4] == b'%PDF', (r.status_code, r.location))
    if shutil.which('pdftotext'):
        pdf = os.path.join(work, 'rapport.pdf')
        open(pdf, 'wb').write(r.data)
        text = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True).stdout
        ok('Synthèse de la comparaison' in text and text.count('Lecture automatique') == 2, 'lecture absente du PDF')
        ok('Hausse liée à la campagne <2024> & suivi.' in text.replace('\n', ' ').replace('  ', ' '), 'commentaire libre absent du PDF')
    print("OK — rapports Word, Excel et PDF : synthèse, lecture automatique et commentaire de l'analyste")

print()
print(f"TOUS LES TESTS D'INTERPRÉTATION MULTI-ENQUÊTE SONT PASSÉS ({checks} vérifications)")
