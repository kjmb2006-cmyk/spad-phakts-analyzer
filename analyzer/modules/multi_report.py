"""
SPAD Analyzer — Rapport d'analyse multi-enquête (PDF, Word, Excel)

Le rapport reprend ce que la page « Analyse multi-enquête » montre, dans l'ordre où on le lit :

  1. les enquêtes comparées, dans l'ordre du temps (Factum), avec leur qualité ;
  2. les contrôles de cohérence entre enquêtes (unités, codages, doublons) ;
  3. la synthèse et un tableau récapitulatif des indicateurs ;
  4. le détail de chaque comparaison, des écarts les plus nets aux indicateurs stables :
     tableaux, graphique, évolution entre enquêtes consécutives, lecture automatique et
     commentaire de l'analyste ;
  5. les notes de méthode (tests, correction des tests multiples, intervalles de confiance).

Le contenu est d'abord construit une seule fois sous forme de blocs neutres
(build_blocks), puis chaque format les met en page : les trois rapports disent donc la
même chose. Les graphiques sont tracés en image (Kaleido) pour les MAX_CHARTS premiers
indicateurs ; au-delà, les tableaux et la lecture suffisent et un message le précise. Le détail
est limité aux MAX_DETAILS premiers indicateurs (le récapitulatif les liste tous).
"""
import io
import math
import os
import re
from datetime import datetime
from xml.sax.saxutils import escape as _esc

import pandas as pd

from modules import multi_comments as mc

MAX_DETAILS = 60             # indicateurs détaillés ; le récapitulatif les liste tous
MAX_CHARTS = 30
MAX_ROWS = 40                # lignes d'un tableau dans le PDF / Word ; l'Excel les garde (jusqu'à 200)

ALIGN_LABELS = {'phakts': 'code PHAKTS complet (Age__1Y et Age__1M sont deux variables)',
                'exact': 'nom exact de colonne'}
SEVERITY = {'erreur': 'Erreur', 'a_verifier': 'À vérifier', 'info': 'Info'}
KIND_LABELS = {'categorical': 'catégorielle', 'continuous': 'continue'}

METHODOLOGY = [
    "Alignement : les colonnes de chaque enquête sont rapprochées par leur code PHAKTS complet ; une même question posée "
    "avec une autre unité (Age__1Y et Age__1M) donne deux variables, jamais comparées ensemble.",
    "Ordre : les enquêtes sont rangées par la date de leur Factum ; à défaut, par ordre d'import.",
    "Comparaison globale : χ² d'indépendance pour une variable catégorielle ; ANOVA si la distribution est normale dans "
    "chaque enquête (Shapiro-Wilk), Kruskal-Wallis sinon, pour une variable continue.",
    "Évolution : entre enquêtes consécutives seulement, avec un intervalle de confiance à 95 % (Newcombe-Wilson pour une "
    "différence de proportions, Welch pour une différence de moyennes). Un écart est net lorsque l'intervalle exclut 0.",
    "Tests multiples : avec de nombreux indicateurs testés à 5 %, certains paraissent significatifs par hasard. Les valeurs de p "
    "sont corrigées (Benjamini-Hochberg) ; q est la proportion attendue de fausses découvertes parmi les indicateurs signalés. "
    "On juge sur q, pas sur p.",
    "Intensité de l'écart : V de Cramér (< 0,1 très faible ; < 0,3 faible ; < 0,5 modérée ; ≥ 0,5 forte). Elle ne dépend pas "
    "des effectifs, contrairement à p.",
    "Limites : un test signale qu'un écart est peu probablement dû au hasard, pas qu'il est important ni qu'il a une cause "
    "précise. Il ne traduit une évolution réelle que si les enquêtes portent sur des populations et des méthodes "
    "d'échantillonnage comparables.",
]


# ─── Mise en forme commune ───────────────────────────────────────────────────

def _fr(x, nd=1):
    return mc._fr(x, nd)


def _p(x):
    return '< 0,0001' if x is not None and x < 0.0001 else _fr(x, 4)


def _cut(df, limit):
    """(df tronqué, message ou '') : les tableaux très longs sont coupés dans le PDF et le Word."""
    if limit and len(df) > limit:
        return df.head(limit), f"{limit} premières lignes sur {len(df)} ; le tableau complet figure dans le rapport Excel."
    return df, ''


def _num(v, nd=None):
    """Nombre à la française : nd décimales fixes, ou le plus court (35 ; 34,83) quand nd est None."""
    if isinstance(v, bool) or not isinstance(v, (int, float)) or (isinstance(v, float) and math.isnan(v)):
        return v
    if nd is not None:
        return _fr(float(v), nd)
    if float(v) == int(v):
        return str(int(v))
    return f'{v:.2f}'.rstrip('0').rstrip('.').replace('.', ',')


def frame_from_index(df, first_column, nd=None):
    """DataFrame dont l'index devient une première colonne nommée ; nombres à la française."""
    out = df.copy().astype(object)
    for col in out.columns:
        out[col] = [_num(v, nd) for v in out[col]]
    out.columns = [str(c) for c in out.columns]
    out.insert(0, first_column, [str(i) for i in out.index])
    return out.reset_index(drop=True)


def evolution_frame(rows):
    out = []
    for e in rows or []:
        pct = e.get('unite') == 'points de %'
        nd, unit = (1, ' %') if pct else (2, '')
        out.append({
            'De → vers': f"{e['de']} → {e['vers']}",
            'Modalité': e['modalite'],
            'Avant': f"{_fr(e['v1'], nd)}{unit} (n={e['n1']})",
            'Après': f"{_fr(e['v2'], nd)}{unit} (n={e['n2']})",
            'Écart': mc._signed(e['diff'], nd) + (' pts' if pct else '') + (' *' if e.get('significatif') else ''),
            'IC 95 %': f"[{_fr(e['ic_bas'], nd)} ; {_fr(e['ic_haut'], nd)}]",
            'p': _p(e.get('p_value')),
        })
    return pd.DataFrame(out)


def recap_frame(ranked_results, corrected):
    out = []
    for r in ranked_results:
        t = r.get('test') or {}
        out.append({
            'Indicateur': r.get('variable') or r.get('radical') or '',
            'Libellé': r.get('label') or '',
            'Type': KIND_LABELS.get(r.get('kind'), ''),
            'Test': t.get('test', '—'),
            'p': _p(t['p_value']) if t else '—',
            'q': _p(t['q_value']) if corrected and t.get('q_value') is not None else '—',
            'V de Cramér': _fr(t['cramers_v'], 2) if t.get('cramers_v') is not None else '—',
            'Verdict': mc.verdict(r, corrected),
        })
    return pd.DataFrame(out)


def pile_frame(pile):
    out = []
    for e in pile:
        q = e.get('quality')
        out.append({'Rang': e['rank'], 'Enquête': e['name'], 'Factum': e.get('factum') or '—',
                    'Obs.': e.get('n_obs', ''), 'Variables': e.get('n_vars', ''),
                    'Source': e.get('source') or 'excel',
                    'Qualité': (f"{_fr(q['pct_sans_anomalie'])} % sans anomalie" if q else '—')})
    return pd.DataFrame(out)


def checks_frame(checks):
    return pd.DataFrame([{'Gravité': SEVERITY.get(c['severite'], c['severite']), 'Contrôle': c['titre'],
                          'Détail': c['detail']} for c in checks])


# ─── Construction du contenu ─────────────────────────────────────────────────

def build_blocks(data, tmp_images):
    """Contenu du rapport en blocs neutres. `tmp_images` reçoit les chemins d'images temporaires créées
    (à supprimer par l'appelant). data : voir app._ms_report_payload()."""
    from modules import report_generator as rg           # tracé des graphiques (Kaleido)

    results = data['results']
    corrected = sum(1 for r in results if r.get('test')) >= 2
    ranked = mc.ranked(results, corrected)
    established = sum(1 for r in ranked if mc.verdict(r, corrected) == 'Écart établi')
    when = data.get('generated_at') or datetime.now()
    order = data.get('order') or []

    blocks = [
        {'t': 'cover', 'title': data.get('title') or "Rapport d'analyse multi-enquête",
         'subtitle': "Analyse multi-enquête (DPF / PHAKTS) — SPAD Analyzer",
         'rows': [('Auteur', data.get('author') or 'Équipe SPAD'),
                  ('Date', when.strftime('%d/%m/%Y à %H:%M')),
                  ('Enquêtes comparées', f"{len(order)} : " + ' → '.join(order)),
                  ('Observations fusionnées', str(data.get('n_obs_merged', '—'))),
                  ('Indicateurs comparés', f"{len(ranked)} dont {established} avec un écart établi"),
                  ('Alignement des variables', ALIGN_LABELS.get(data.get('mode'), str(data.get('mode'))))]},
        {'t': 'pagebreak'},
        {'t': 'h1', 'text': '1. Enquêtes comparées'},
        {'t': 'p', 'text': "Dans l'ordre du temps (date du Factum). La qualité est la part des lignes sans anomalie aux contrôles "
                           "de validation des données (doublons, heures, seuils, valeurs extrêmes, règles propres à l'enquête)."},
        {'t': 'table', 'df': pile_frame(data['pile']), 'widths': [1.6, 3.5, 5.3, 1.3, 2.2, 1.9, 3.1]},
    ]

    checks = data.get('checks') or []
    blocks.append({'t': 'h1', 'text': '2. Contrôles de cohérence entre enquêtes'})
    if checks:
        blocks += [{'t': 'p', 'text': "À lever avant d'interpréter les écarts : ils peuvent venir des données elles-mêmes "
                                      "(unités, codages des réponses, enquête ajoutée deux fois)."},
                   {'t': 'table', 'df': checks_frame(checks), 'widths': [1.8, 4.5, 11.3]}]
    else:
        blocks.append({'t': 'p', 'text': "Aucun problème de cohérence détecté entre les enquêtes (unités, codages, doublons, Factum)."})

    blocks.append({'t': 'h1', 'text': '3. Synthèse'})
    if data.get('synthese'):
        blocks.append({'t': 'text', 'text': data['synthese']})
    if len(ranked):
        blocks += [{'t': 'h2', 'text': 'Récapitulatif des indicateurs'},
                   {'t': 'p', 'text': "Classés des écarts les plus nets aux indicateurs stables."
                                      + (" q : valeur de p corrigée des tests multiples (Benjamini-Hochberg)." if corrected else '')},
                   {'t': 'table', 'df': recap_frame(ranked, corrected), 'limit': 80,
                    'widths': [2.8, 2.4, 2.5, 3.0, 1.7, 1.7, 1.6, 2.2]}]

    blocks += [{'t': 'sheet', 'name': 'Comparaisons'}, {'t': 'pagebreak'},
               {'t': 'h1', 'text': '4. Comparaisons détaillées'}]
    if len(ranked) > MAX_DETAILS:
        blocks.append({'t': 'p', 'style': 'note',
                       'text': f"Détail des {MAX_DETAILS} premiers indicateurs sur {len(ranked)} (les écarts les plus nets d'abord) ; "
                               "le récapitulatif de la section 3 les liste tous et la page d'analyse montre le détail de chacun."})
    charts = 0
    for i, r in enumerate(ranked[:MAX_DETAILS], start=1):
        variable = r.get('variable') or r.get('radical') or '?'
        blocks.append({'t': 'h2', 'text': f"4.{i}  {variable}" + (f" — {r['label']}" if r.get('label') else '')})
        if r.get('error'):
            blocks.append({'t': 'p', 'text': f"Calcul impossible : {r['error']}"})
            continue
        t = r.get('test')
        head = f"Variable {KIND_LABELS.get(r.get('kind'), '')}"
        if t:
            head += (f" · {t['test']} : p {'< 0,0001' if t['p_value'] < 0.0001 else '= ' + _fr(t['p_value'], 4)}"
                     + (f", q {'< 0,0001' if t['q_value'] < 0.0001 else '= ' + _fr(t['q_value'], 4)}"
                        if corrected and t.get('q_value') is not None else ''))
        blocks.append({'t': 'p', 'text': f"{head} · Verdict : {mc.verdict(r, corrected)}", 'style': 'meta'})

        frames = r.get('frames') or {}
        if r.get('kind') == 'categorical' and 'counts' in frames:
            blocks += [{'t': 'table', 'title': 'Effectifs', 'df': frame_from_index(frames['counts'], 'Modalité'), 'limit': MAX_ROWS},
                       {'t': 'table', 'title': 'Pourcentages (% colonne)', 'df': frame_from_index(frames['pct'], 'Modalité', nd=1),
                        'limit': MAX_ROWS}]
        elif 'summary' in frames:
            blocks.append({'t': 'table', 'title': 'Résumé statistique par enquête',
                           'df': frame_from_index(frames['summary'], 'Enquête')})
        if r.get('chart'):
            if charts < MAX_CHARTS:
                path = rg._fig_to_img(r['chart'], width=640, height=380)
                if path:
                    tmp_images.append(path)
                    charts += 1
                    blocks.append({'t': 'image', 'path': path, 'caption': f"Figure : comparaison de {variable} entre enquêtes"})
            elif charts == MAX_CHARTS:
                charts += 1
                blocks.append({'t': 'p', 'style': 'note',
                               'text': f"Les graphiques ne sont tracés que pour les {MAX_CHARTS} premiers indicateurs (ceux aux écarts "
                                       "les plus nets) ; la page d'analyse les montre tous."})
        evo = evolution_frame(r.get('evolution'))
        if len(evo):
            blocks += [{'t': 'table', 'title': 'Évolution entre enquêtes consécutives', 'df': evo, 'limit': MAX_ROWS,
                        'widths': [3.6, 3.4, 2.5, 2.5, 2.0, 2.0, 1.6]},
                       {'t': 'p', 'style': 'note', 'text': "* l'intervalle de confiance à 95 % exclut 0 : écart net entre les deux enquêtes."}]
        blocks.append({'t': 'lecture', 'auto': r.get('interpretation') or '', 'user': r.get('user_comment') or ''})

    blocks += [{'t': 'sheet', 'name': 'Méthode'}, {'t': 'h1', 'text': '5. Notes de méthode'}]
    blocks += [{'t': 'bullet', 'text': line} for line in METHODOLOGY]
    return blocks


# ─── PDF ─────────────────────────────────────────────────────────────────────

def _pdf_text(text):
    return _esc(str(text)).replace('\n', '<br/>')


def render_pdf(blocks, data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)
    from modules import report_generator as rg

    S = rg._styles()
    S['Note'] = ParagraphStyle('Note', fontSize=8, textColor=rg.GRAY, spaceAfter=6, leading=10)
    S['Auto'] = ParagraphStyle('Auto', fontSize=9, textColor=rg.PRIMARY, backColor=rg.LIGHT, borderPad=6,
                               spaceAfter=4, leading=12.5, leftIndent=8, rightIndent=8)
    S['User'] = ParagraphStyle('User', fontSize=9.5, textColor=colors.HexColor('#2C3E50'), spaceAfter=8, leading=13,
                               leftIndent=8, rightIndent=8, fontName='Helvetica-Oblique')
    S['Cell'] = ParagraphStyle('Cell', fontSize=7.5, leading=9.2)
    S['CellHead'] = ParagraphStyle('CellHead', fontSize=7.5, leading=9.2, textColor=colors.white, fontName='Helvetica-Bold')
    S['Bullet'] = ParagraphStyle('Bullet', fontSize=9.5, leading=13, leftIndent=12, bulletIndent=2, spaceAfter=3)
    usable = 17 * cm

    def table(df, widths=None):
        cols = [str(c) for c in df.columns]
        if widths and len(widths) == len(cols):
            total = sum(widths)
            col_w = [w / total * usable for w in widths]
        else:
            weight = [max(4, min(34, max([len(c)] + [len(str(v)) for v in df[c].head(30)]))) for c in df.columns]
            col_w = [w / sum(weight) * usable for w in weight]
        rows = [[Paragraph(_esc(c), S['CellHead']) for c in cols]]
        rows += [[Paragraph(_esc(str(v)), S['Cell']) for v in row] for row in df.astype(object).where(df.notna(), '').values.tolist()]
        t = Table(rows, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), rg.PRIMARY), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [rg.WHITE, rg.LIGHT]),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#dee2e6')),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3), ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ]))
        return t

    story = []
    for b in blocks:
        kind = b['t']
        if kind == 'cover':
            meta = Table([[Paragraph(_esc(k), S['Body']), Paragraph(_esc(v), S['Body'])] for k, v in b['rows']],
                         colWidths=[5 * cm, 12 * cm])
            meta.setStyle(TableStyle([('BACKGROUND', (0, 0), (0, -1), rg.LIGHT), ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dee2e6')),
                                      ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                                      ('LEFTPADDING', (0, 0), (-1, -1), 8), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story += [Spacer(1, 3 * cm), Paragraph(_esc(b['title']), S['Title']), Spacer(1, 0.3 * cm),
                      Paragraph(_esc(b['subtitle']), S['Subtitle']), HRFlowable(width='90%', thickness=2, color=rg.PRIMARY, spaceAfter=12),
                      Spacer(1, 0.4 * cm), meta]
        elif kind == 'pagebreak':
            story.append(PageBreak())
        elif kind == 'h1':
            story += [Paragraph(_esc(b['text']), S['H1']), HRFlowable(width='100%', thickness=1, color=rg.ACCENT, spaceAfter=8)]
        elif kind == 'h2':
            story.append(Paragraph(_esc(b['text']), S['H2']))
        elif kind == 'p':
            story.append(Paragraph(_pdf_text(b['text']), S['Note' if b.get('style') == 'note' else 'Meta' if b.get('style') == 'meta' else 'Body']))
        elif kind == 'text':
            story.append(Paragraph(_pdf_text(b['text']), S['Auto']))
        elif kind == 'bullet':
            story.append(Paragraph(_esc(b['text']), S['Bullet'], bulletText='•'))
        elif kind == 'table':
            df, cut_note = _cut(b['df'], b.get('limit'))
            if b.get('title'):
                story.append(Paragraph(f"<b>{_esc(b['title'])}</b>", S['Body']))
            story.append(table(df, b.get('widths')))
            story.append(Spacer(1, 0.15 * cm))
            if cut_note:
                story.append(Paragraph(_esc(cut_note), S['Note']))
        elif kind == 'image':
            story += [Spacer(1, 0.2 * cm), Image(b['path'], width=15 * cm, height=15 * cm * 380 / 640),
                      Paragraph(_esc(b['caption']), S['Caption'])]
        elif kind == 'lecture':
            if b['auto']:
                story.append(Paragraph('<b>Lecture automatique</b><br/>' + _pdf_text(b['auto']), S['Auto']))
            if b['user']:
                story.append(Paragraph("<b>Commentaire de l'analyste :</b> " + _pdf_text(b['user']), S['User']))
            story.append(Spacer(1, 0.3 * cm))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, topMargin=2.5 * cm, bottomMargin=2.2 * cm, leftMargin=2 * cm, rightMargin=2 * cm,
                      title=data.get('title') or 'Rapport multi-enquête').build(story, canvasmaker=rg._NumberedCanvas)
    return buf.getvalue()


# ─── Word ────────────────────────────────────────────────────────────────────

def render_word(blocks, data):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor

    doc = Document()
    doc.styles['Normal'].font.size = Pt(10)
    for name, size, color in (('Title', 22, (26, 60, 94)), ('Heading 1', 15, (26, 60, 94)), ('Heading 2', 12, (224, 92, 26))):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(*color)

    def shade(cell, hex_fill):
        pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_fill)
        pr.append(shd)

    def table(df, widths=None):
        t = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
        t.style = 'Table Grid'
        if widths and len(widths) == len(df.columns):
            t.autofit = False
            col_w = [Cm(16 * w / sum(widths)) for w in widths]
            for j, column in enumerate(t.columns):
                column.width = col_w[j]
            for row in t.rows:
                for j, cell in enumerate(row.cells):
                    cell.width = col_w[j]
        header = OxmlElement('w:tblHeader')               # l'en-tête se répète en haut de chaque page
        header.set(qn('w:val'), 'true')
        t.rows[0]._tr.get_or_add_trPr().append(header)
        for j, col in enumerate(df.columns):
            cell = t.rows[0].cells[j]
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(col))
            run.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(255, 255, 255)
            shade(cell, '1A3C5E')
        for i, row in enumerate(df.astype(object).where(df.notna(), '').values.tolist(), start=1):
            for j, v in enumerate(row):
                cell = t.rows[i].cells[j]
                cell.text = ''
                cell.paragraphs[0].add_run(str(v)).font.size = Pt(8)
        doc.add_paragraph()

    for b in blocks:
        kind = b['t']
        if kind == 'cover':
            doc.add_paragraph(b['title'], style='Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub = doc.add_paragraph(b['subtitle'])
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
            t = doc.add_table(rows=len(b['rows']), cols=2)
            t.style = 'Table Grid'
            for i, (k, v) in enumerate(b['rows']):
                t.rows[i].cells[0].text = k
                t.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
                shade(t.rows[i].cells[0], 'F0F6FC')
                t.rows[i].cells[1].text = v
        elif kind == 'pagebreak':
            doc.add_page_break()
        elif kind == 'h1':
            doc.add_heading(b['text'], level=1)
        elif kind == 'h2':
            doc.add_heading(b['text'], level=2)
        elif kind in ('p', 'text'):
            p = doc.add_paragraph()
            run = p.add_run(b['text'])
            if b.get('style') in ('note', 'meta'):
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(108, 117, 125)
        elif kind == 'bullet':
            doc.add_paragraph(b['text'], style='List Bullet')
        elif kind == 'table':
            df, cut_note = _cut(b['df'], b.get('limit'))
            if b.get('title'):
                doc.add_paragraph().add_run(b['title']).bold = True
            table(df, b.get('widths'))
            if cut_note:
                doc.add_paragraph().add_run(cut_note).italic = True
        elif kind == 'image':
            doc.add_picture(b['path'], width=Inches(6.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(b['caption'])
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(8.5)
        elif kind == 'lecture':
            if b['auto']:
                p = doc.add_paragraph()
                p.add_run('Lecture automatique\n').bold = True
                p.add_run(b['auto'])
            if b['user']:
                p = doc.add_paragraph()
                p.add_run("Commentaire de l'analyste : ").bold = True
                p.add_run(b['user']).italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─── Excel ───────────────────────────────────────────────────────────────────

def render_excel(blocks, data):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    HEAD_FILL = PatternFill('solid', fgColor='1A3C5E')
    HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
    BODY = Font(size=10)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Résumé'
    state = {'row': 1}
    WIDE = 10                      # colonnes fusionnées pour les textes

    def put(text, font=BODY, wrap=False, fill=None):
        r = state['row']
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        if fill:
            cell.fill = fill
        if wrap:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=WIDE)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            lines = sum(max(1, math.ceil(len(part) / 120)) for part in str(text).split('\n'))
            ws.row_dimensions[r].height = 14 * lines + 2
        state['row'] = r + 1

    def widen(j, width):
        letter = get_column_letter(j)
        current = ws.column_dimensions[letter].width or 8
        ws.column_dimensions[letter].width = max(current, min(width, 48))

    for b in blocks:
        kind = b['t']
        if kind == 'sheet':
            ws = wb.create_sheet(re.sub(r'[\[\]:*?/\\]', '', b['name'])[:31])
            state['row'] = 1
        elif kind == 'cover':
            put(b['title'], Font(color='1A3C5E', bold=True, size=16))
            put(b['subtitle'], Font(size=10, italic=True))
            state['row'] += 1
            for k, v in b['rows']:
                r = state['row']
                ws.cell(row=r, column=1, value=k).font = Font(size=10, bold=True)
                ws.cell(row=r, column=2, value=v).font = BODY
                state['row'] += 1
            widen(1, 28)
            widen(2, 48)
        elif kind == 'h1':
            state['row'] += 1
            put(b['text'], Font(color='1A3C5E', bold=True, size=13))
        elif kind == 'h2':
            state['row'] += 1
            put(b['text'], Font(color='E05C1A', bold=True, size=11))
        elif kind in ('p', 'text', 'bullet'):
            put(('• ' if kind == 'bullet' else '') + b['text'], Font(size=9, italic=True) if b.get('style') in ('note', 'meta') else BODY,
                wrap=True)
        elif kind == 'table':
            df, _ = _cut(b['df'], 200 if b.get('limit') else None)
            if b.get('title'):
                put(b['title'], Font(size=10, bold=True))
            r = state['row']
            for j, col in enumerate(df.columns, start=1):
                c = ws.cell(row=r, column=j, value=str(col))
                c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, Alignment(horizontal='center', vertical='center', wrap_text=True)
            for i, row in enumerate(df.astype(object).where(df.notna(), '').values.tolist(), start=1):
                for j, v in enumerate(row, start=1):
                    c = ws.cell(row=r + i, column=j, value=v)
                    c.font = BODY
                    c.alignment = Alignment(vertical='top', wrap_text=len(str(v)) > 40)
            for j, col in enumerate(df.columns, start=1):
                widen(j, max([len(str(col))] + [min(len(str(v)), 60) for v in df[col].head(60)]) + 2)
            state['row'] = r + len(df) + 2
        elif kind == 'image':
            try:
                img = XLImage(b['path'])
                scale = 560 / float(img.width)
                img.width, img.height = 560, int(img.height * scale)
                ws.add_image(img, f"A{state['row']}")
                state['row'] += max(2, int(img.height / 15) + 1)
            except Exception:
                pass
        elif kind == 'lecture':
            if b['auto']:
                put('Lecture automatique', Font(size=10, bold=True))
                put(b['auto'], BODY, wrap=True)
            if b['user']:
                put("Commentaire de l'analyste : " + b['user'], Font(size=10, italic=True), wrap=True)
            state['row'] += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


RENDERERS = {'pdf': (render_pdf, 'application/pdf'),
             'docx': (render_word, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
             'xlsx': (render_excel, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}


def generate(data, fmt):
    """Rapport complet : (octets, type MIME). Les images temporaires sont supprimées avant de rendre la main."""
    render, mime = RENDERERS[fmt]
    tmp = []
    try:
        return render(build_blocks(data, tmp), data), mime
    finally:
        for path in tmp:
            try:
                os.remove(path)
            except OSError:
                pass
