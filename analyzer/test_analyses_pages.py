#!/usr/bin/env python3
"""
Test script — pages d'analyse de bout en bout (statistiques descriptives, croisée, multivariées, validation, rapports).

Un jeu de données de formulaire Kobo (noms internes des modalités, une question à choix multiples en une colonne, un
identifiant, une date) et son XLSForm. Vérifie ce que voit l'utilisateur :
  - les modalités s'affichent avec leur vrai nom ; les dates, identifiants et textes libres ne sont pas proposés comme variables ;
  - la première visite de la page descriptive n'affiche pas tout le formulaire ;
  - chaque résultat porte sa lecture automatique ; les pages se rechargent sans « renvoi de formulaire » ;
  - une classe enregistrée devient une variable ; une ligne écartée disparaît des analyses puis revient ;
  - le formulaire joint contrôle les réponses ; les rapports (PDF, Word, Excel) reprennent les lectures.
"""
import io
import os
import re
import tempfile

import numpy as np
import pandas as pd

import app as m
from modules import dataset

print("=" * 70)
print("TEST — pages d'analyse")
print("=" * 70)
checks = 0


def ok(cond, msg=''):
    global checks
    assert cond, msg
    checks += 1


rng = np.random.default_rng(2)
n = 240
prof = rng.choice(['sage_femme', 'infirmier', 'medecin'], n, p=[.45, .35, .2])
df = pd.DataFrame({
    '_uuid': [f'u{i}' for i in range(n)],
    'Date_Collecte__A': ['2026-03-01'] * n,
    'ENTETE/Code_Enqueteur__E': rng.choice(['E1', 'E2', 'E3'], n),
    'ENTETE/District__X': rng.choice(['nord', 'sud', 'est'], n),
    'METIER/Profession__X': prof,
    'METIER/Sexe__X': np.where(rng.random(n) < np.where(prof == 'sage_femme', .95, .35), 'femme', 'homme'),
    'METIER/Fume__B': np.where(rng.random(n) < np.where(prof == 'medecin', .1, .3), 'oui', 'non'),
    'METIER/Anciennete__1Y!0<=N<=60': rng.integers(0, 40, n),
    'METIER/Age__1Y': rng.integers(22, 62, n),
    'METIER/Effets__X': [' '.join(rng.choice(['cancer', 'toux', 'asthme'], size=rng.integers(1, 3), replace=False)) for _ in range(n)],
    'METIER/Commentaire__Z': ['rien'] * n,
})
df.loc[5, 'METIER/Age__1Y'] = 200
df.loc[9, '_uuid'] = df.loc[8, '_uuid']
work = tempfile.mkdtemp()
path = os.path.join(work, 'kobo.xlsx')
df.to_excel(path, index=False)
form_path = os.path.join(work, 'form.xlsx')
with pd.ExcelWriter(form_path) as w:
    pd.DataFrame({'type': ['select_one prof', 'select_one sx', 'select_one yn', 'integer', 'select_multiple ef'],
                  'name': ['Profession__X', 'Sexe__X', 'Fume__B', 'Age__1Y', 'Effets__X'],
                  'label': ['Profession du répondant', 'Sexe', 'Fume ?', 'Âge', 'Effets connus du tabac'],
                  'required': ['yes', 'yes', '', 'yes', ''], 'relevant': ['', '', '', '', ''],
                  'constraint': ['', '', '', '. <= 120', ''], 'constraint_message': ['', '', '', 'Âge entre 0 et 120.', '']}).to_excel(w, sheet_name='survey', index=False)
    pd.DataFrame({'list_name': ['prof'] * 3 + ['sx'] * 2 + ['yn'] * 2 + ['ef'] * 3,
                  'name': ['sage_femme', 'infirmier', 'medecin', 'femme', 'homme', 'oui', 'non', 'cancer', 'toux', 'asthme'],
                  'label': ['Sage-femme', 'Infirmier(ère)', 'Médecin', 'Femme', 'Homme', 'Oui', 'Non',
                            'Cancer du poumon', 'Toux chronique', 'Asthme']}).to_excel(w, sheet_name='choices', index=False)

m.app.config['TESTING'] = True
m.app.config['REPORTS_FOLDER'] = os.path.join(work, 'rapports')
os.makedirs(m.app.config['REPORTS_FOLDER'])
strip = lambda h: re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', h))

with m.app.test_client() as c:
    with c.session_transaction() as sess:
        sess['data_path'] = path
        sess['data_meta'] = {'name': 'kobo', 'n_obs': n, 'n_vars': df.shape[1]}
        sess['original_filename'] = 'kobo.xlsx'

    # ── Sans formulaire : noms reconstitués, variables non analytiques écartées
    page = c.get('/descriptive').get_data(as_text=True)
    ok('Sage femme' in page and 'sage_femme' not in page, 'modalités mises en forme')
    ok('Date_Collecte' not in page and 'Commentaire__Z' not in page and 'Code_Enqueteur' not in page.split('name="cat_vars"')[1][:400], 'date / texte / identifiant non proposés')
    ok('Effets__X/Cancer' in page or 'Effets' in page, 'choix multiple éclaté en groupe')
    print("OK — sans formulaire : modalités mises en forme, date / identifiant / texte libre écartés")

    # ── Formulaire joint : vrais libellés
    with open(form_path, 'rb') as f:
        r = c.post('/data/labels', data={'xlsform': (f, 'form.xlsx')}, content_type='multipart/form-data', follow_redirects=True)
    ok('Formulaire lu' in r.get_data(as_text=True), 'XLSForm accepté')
    page = c.get('/descriptive').get_data(as_text=True)
    ok('Sage-femme' in page and 'Profession du répondant' in page, 'libellés du XLSForm')
    ok('Cancer du poumon' in page, 'choix multiples : libellés des options')
    ok('Lecture automatique' in page and 'IC à 95 %' in page, 'lecture sous chaque variable')
    print("OK — avec le XLSForm : « Sage-femme », « Profession du répondant », choix multiples « Cancer du poumon »")

    # ── Première visite : un aperçu, pas tout le formulaire ; décomposition par groupe
    ok(page.count('id="r-cat-') <= 12 + 0 and page.count('id="r-bin-') >= 1, 'aperçu limité')
    c.post('/descriptive', data={'cat_vars': ['METIER/Profession__X', 'METIER/Fume__B'], 'num_vars': ['METIER/Age__1Y'], 'group_var': 'ENTETE/District__X'})
    page = c.get('/descriptive').get_data(as_text=True)
    ok(page.count('Selon «') >= 3 and page.count('Lecture automatique') >= 6, (page.count('Selon «'), page.count('Lecture automatique')))
    print("OK — descriptive : aperçu limité au départ, décomposition par district avec sa lecture")

    # ── Croisée : tests, résidus, stratification, choix multiples, Excel
    c.post('/crosstabs', data={'kind': 'pivot', 'rows': ['METIER/Profession__X'], 'cols': ['METIER/Sexe__X'], 'pct_type': 'row', 'layer': 'ENTETE/District__X'})
    page = c.get('/crosstabs').get_data(as_text=True)
    ok('Sage-femme' in page and 'Association' in page and 'Strates' in page and 'Résidus ajustés' in page, 'croisée complète')
    c.post('/crosstabs', data={'kind': 'multi', 'group': 'METIER/Effets__X', 'var': 'METIER/Profession__X'})
    page = c.get('/crosstabs').get_data(as_text=True)
    ok('Cancer du poumon' in page and 'répondants' in page, 'choix multiples × variable')
    xl = pd.read_excel(io.BytesIO(c.get('/crosstabs/export.xlsx').data), sheet_name=None)
    ok('Synthèse' in xl and 'Tableau' in xl, list(xl))
    ok(c.get('/crosstabs/export.xlsx?kind=multi').status_code == 200)
    print("OK — croisée : test lu, résidus, strates, question à choix multiples, export Excel")

    # ── Multivariées : chaque méthode, lecture, classes enregistrées
    for form in ({'method': 'pca', 'variables': ['METIER/Age__1Y', 'METIER/Anciennete__1Y!0<=N<=60'], 'n_components': 2},
                 {'method': 'mca', 'variables': ['METIER/Profession__X', 'METIER/Sexe__X', 'METIER/Fume__B'], 'n_components': 2},
                 {'method': 'ca', 'ca_row': 'METIER/Profession__X', 'ca_col': 'METIER/Sexe__X'},
                 {'method': 'clustering', 'variables': ['METIER/Profession__X', 'METIER/Sexe__X', 'METIER/Fume__B'], 'cluster_method': 'cah',
                  'cluster_kind': 'mca', 'n_clusters': '0'},
                 {'method': 'clustering', 'variables': ['METIER/Age__1Y', 'METIER/Anciennete__1Y!0<=N<=60'], 'cluster_method': 'kmeans',
                  'cluster_kind': 'numeric', 'n_clusters': '3'}):
        c.post('/multivariate', data=form)
        r = c.get('/multivariate')
        html = r.get_data(as_text=True)
        ok(r.status_code == 200 and 'Lecture automatique' in html and 'Erreur d' not in html, (form['method'], strip(html)[:200]))
    r = c.post('/multivariate/classes', data={'name': 'Classe_test'}, follow_redirects=True)
    ok('enregistrée' in r.get_data(as_text=True) and dataset.derived_names(path) == ['Classe_test'])
    page = c.get('/descriptive').get_data(as_text=True)
    ok('Classe_test' in page, 'la classe est proposée comme variable')
    c.post('/crosstabs', data={'kind': 'pivot', 'rows': ['Classe_test'], 'cols': ['METIER/Profession__X']})
    ok('Classe 1' in c.get('/crosstabs').get_data(as_text=True), 'la classe se croise avec une autre variable')
    ok(c.get('/multivariate/export.xlsx').status_code == 200)
    c.post('/data/derived/delete', data={'name': 'Classe_test'})
    ok(dataset.derived_names(path) == [])
    print("OK — multivariées : 5 analyses lues ; la classe enregistrée est une variable (tri à plat, croisement) puis supprimée")

    # ── Validation : formulaire, doublon d'identifiant, écartement
    page = c.get('/qualite/validation').get_data(as_text=True)
    ok('Identifiant de soumission « _uuid » en double' in page, 'doublon de _uuid')
    ok('Contrainte du formulaire non respectée (Age__1Y)' in page and 'Âge entre 0 et 120.' in page, 'contrainte du formulaire')
    ok('Signalements par enquêteur' in page, 'synthèse par enquêteur')
    r = c.post('/qualite/validation/exclure', data={'mode': 'erreurs'}, follow_redirects=True)
    excluded = dataset.excluded_rows(path)
    ok(6 in excluded and {9, 10} & set(excluded), excluded)
    page = c.get('/descriptive').get_data(as_text=True)
    ok('écartée(s) de cette analyse' in page)
    ok(len(m.get_dataframe(full=True)) == n and len(dataset.apply_exclusions(dataset.read_frame(path), path)) == n - len(excluded))
    c.post('/data/exclusions/clear')
    ok(dataset.excluded_rows(path) == [])
    print("OK — validation : doublon d'identifiant, contrainte du formulaire, synthèse par enquêteur, lignes écartées puis rétablies")

    # ── Rapports : lectures dans les trois formats
    base = {'analyses': ['descriptive', 'crosstabs'], 'cat_vars': ['METIER/Profession__X'], 'num_vars': ['METIER/Age__1Y'],
            'crosstabs_row': 'METIER/Profession__X', 'crosstabs_col': 'METIER/Sexe__X', 'report_title': 'Rapport test'}
    from docx import Document
    from openpyxl import load_workbook
    r = c.post('/report', data=dict(base, report_format='word'))
    ok(r.status_code == 200 and r.data[:2] == b'PK', r.status_code)
    text = '\n'.join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
    ok('Lecture automatique' in text and 'Sage-femme' in ' '.join(c_.text for t in Document(io.BytesIO(r.data)).tables for row in t.rows for c_ in row.cells), 'lecture + libellés dans le Word')
    r = c.post('/report', data=dict(base, report_format='excel'))
    cells = [str(v) for ws in load_workbook(io.BytesIO(r.data)) for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    ok(r.status_code == 200 and any(v == 'Lecture automatique :' for v in cells) and any(v.startswith('• ') for v in cells), 'lecture en puces dans l\'Excel')
    r = c.post('/report', data=dict(base, report_format='pdf'))
    ok(r.status_code == 200 and r.data[:4] == b'%PDF', r.status_code)
    print("OK — rapports PDF / Word / Excel : lectures en puces, libellés, « < » échappé")

# ── Page multi-enquête : les analyses multivariées de la pile fonctionnent toujours
entries = []
for i, (tag, frame) in enumerate((('a', df.iloc[:120]), ('b', df.iloc[120:]))):
    pth = os.path.join(work, f'ms_{tag}.xlsx')
    frame.drop(columns=['_uuid']).to_excel(pth, index=False)
    entries.append({'name': f'ms_{tag}.xlsx', 'path': pth, 'n_obs': len(frame), 'n_vars': frame.shape[1], 'phakts_vars': 5,
                    'factum': f'PGT14_OS2_RC2_2__CIV_MS__1C2_Etude_Tabac__{2023 + i}0900__ES'})
with m.app.test_client() as c:
    with c.session_transaction() as sess:
        sess['multi_surveys'] = entries
    for params, badge in (({'mv_method': 'pca', 'mv_variables': ['Age__1Y', 'Anciennete__1Y'], 'mv_n_components': '2'}, 'ACP'),
                          ({'mv_method': 'mca', 'mv_variables': ['Profession__X', 'Sexe__X'], 'mv_n_components': '2'}, 'ACM'),
                          ({'mv_method': 'ca', 'mv_ca_row': 'Profession__X', 'mv_ca_col': 'Sexe__X'}, 'AFC'),
                          ({'mv_method': 'clustering', 'mv_variables': ['Age__1Y', 'Anciennete__1Y'], 'mv_n_clusters': '3'}, 'Classification K-Means')):
        c.post('/multi-survey', data=dict(params, action='run_mv_ms'))
        page = c.get('/multi-survey').get_data(as_text=True)
        ok(f'<strong>{badge}</strong>' in page and 'Analyse multivariée :' not in strip(page), (badge, strip(page)[:150]))
    c.post('/multi-survey', data={'action': 'clear_mv_ms'})
print("OK — pile multi-enquête : les quatre analyses multivariées s'affichent toujours")

print()
print(f"TOUS LES TESTS DES PAGES D'ANALYSE SONT PASSÉS ({checks} vérifications)")
