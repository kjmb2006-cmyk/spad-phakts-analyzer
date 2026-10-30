#!/usr/bin/env python3
"""
Test script — rapport d'analyse multi-enquête (PDF, Word, Excel).

Depuis la page « Analyse multi-enquête », trois boutons produisent un rapport : enquêtes
comparées, contrôles de cohérence, synthèse, détail de chaque comparaison (tableaux,
graphique, évolution, lecture automatique, commentaire de l'analyste) et notes de méthode.
Vérifie :
  - le contenu des trois formats, y compris pour un commentaire contenant « < » et « & » ;
  - l'ordre : écarts établis d'abord, indicateurs stables ensuite ;
  - le mode automatique (indicateurs détectés) comme le panier manuel ;
  - les refus propres (moins de deux enquêtes, rien à rapporter) ;
  - le plafond de détail et la mise en forme à la française des tableaux.
"""
import io
import os
import shutil
import subprocess
import tempfile

import numpy as np
import pandas as pd

import app as m
from modules import multi_comments as mc
from modules import multi_report as mr
from modules import report_generator as rg

print("=" * 70)
print("TEST — rapport d'analyse multi-enquête")
print("=" * 70)

checks = 0


def ok(cond, msg=''):
    global checks
    assert cond, msg
    checks += 1


rng = np.random.default_rng(11)
regions = [f'Région {chr(65 + i)}' for i in range(15)]


def survey(oui, age, n, tag, extra=None):
    d = pd.DataFrame({'_uuid': [f'{tag}{i}' for i in range(n)],
                      'Fumeur__B': np.where(rng.random(n) < oui, 'OUI', 'NON'),
                      'Sexe__X': np.where(rng.random(n) < .5, 'Femme', 'Homme'),
                      'Region__X': rng.choice(regions, n),
                      'Age__1Y': rng.normal(age, 6, n).round(0)})
    if extra:
        d[extra] = rng.integers(1, 12, n)       # Age__1M face à Age__1Y : contrôle de cohérence en erreur
    return d


FAC = [f'PGT14_OS2_RC2_2__CIV_MS__1C2_Etude_Tabac__{d}__ES' for d in ('20230900', '20240900', '20250900')]
work = tempfile.mkdtemp()
entries = []
for (name, df), fac in zip((('Enquête 2023', survey(.2, 35, 200, 'a')), ('Enquête 2024', survey(.4, 40, 25, 'b')),
                            ('Enquête 2025', survey(.6, 45, 200, 'c', 'Age__1M'))), FAC):
    path = os.path.join(work, name.replace(' ', '_') + '.xlsx')
    df.to_excel(path, index=False)
    entries.append({'name': os.path.basename(path), 'path': path, 'n_obs': len(df), 'n_vars': df.shape[1],
                    'phakts_vars': df.shape[1], 'factum': fac})

m.app.config['TESTING'] = True
COMMENT = "Hausse cohérente avec la campagne de 2024 <à confirmer> & suivi."
charts_available = rg._fig_to_img(__import__('plotly.express', fromlist=['bar']).bar(x=['a'], y=[1]).to_json(), 200, 150) is not None

# ── Mise en forme à la française, classement
frame = mr.frame_from_index(pd.DataFrame({'E1': [75.5, 24.5], 'E2': [64.0, 36.0]}, index=['NON', 'OUI']), 'Modalité', nd=1)
ok(frame.iloc[0].tolist() == ['NON', '75,5', '64,0'], frame.iloc[0].tolist())
summary = mr.frame_from_index(pd.DataFrame({'Effectif': [200], 'Moyenne': [34.83], 'Médiane': [35.0]}, index=['E1']), 'Enquête')
ok(summary.iloc[0].tolist() == ['E1', '200', '34,83', '35'], summary.iloc[0].tolist())
evo = mr.evolution_frame([{'de': 'A', 'vers': 'B', 'modalite': 'OUI', 'v1': 20.0, 'v2': 43.0, 'n1': 200, 'n2': 25, 'diff': 23.0,
                           'ic_bas': 14.0, 'ic_haut': 31.5, 'p_value': 0.00001, 'significatif': True, 'unite': 'points de %'}])
ok(evo.iloc[0]['Écart'] == '+23,0 pts *' and evo.iloc[0]['IC 95 %'] == '[14,0 ; 31,5]' and evo.iloc[0]['p'] == '< 0,0001', evo.iloc[0].to_dict())
rows = [{'variable': 'stable', 'kind': 'categorical', 'test': {'p_value': .6, 'q_value': .6, 'cramers_v': .05}},
        {'variable': 'tentative', 'kind': 'categorical', 'test': {'p_value': .03, 'q_value': .2, 'cramers_v': .1}},
        {'variable': 'weak', 'kind': 'categorical', 'test': {'p_value': .001, 'q_value': .004, 'cramers_v': .15}},
        {'variable': 'strong', 'kind': 'categorical', 'test': {'p_value': .0001, 'q_value': .0004, 'cramers_v': .6}},
        {'variable': 'untested', 'kind': 'categorical', 'test': None}, {'variable': 'boom', 'error': 'x'}]
ok([r['variable'] for r in mc.ranked(rows)] == ['strong', 'weak', 'tentative', 'stable', 'untested', 'boom'],
   [r['variable'] for r in mc.ranked(rows)])
print("OK — tableaux à la française, évolution chiffrée, classement des écarts les plus nets aux indicateurs stables")

# ── Plafond de détail : le récapitulatif liste tout, le détail s'arrête à MAX_DETAILS
base = {'variable': 'V', 'kind': 'categorical', 'stats': {'surveys': []}, 'evolution': [], 'chart': None,
        'test': {'test': 'χ² (indépendance)', 'p_value': .5, 'q_value': .5, 'cramers_v': .1},
        'frames': {'counts': pd.DataFrame({'E1': [1]}, index=['a']), 'pct': pd.DataFrame({'E1': [100.0]}, index=['a'])}}
many = [dict(base, variable=f'V{i}') for i in range(mr.MAX_DETAILS + 15)]
blocks = mr.build_blocks({'title': 'T', 'author': 'A', 'mode': 'phakts', 'order': ['E1', 'E2'], 'pile': [], 'checks': [],
                          'results': many, 'synthese': 's', 'n_obs_merged': 10}, [])
ok(sum(1 for b in blocks if b['t'] == 'h2' and b['text'].startswith('4.')) == mr.MAX_DETAILS)
recap = next(b for b in blocks if b['t'] == 'table' and 'Verdict' in b['df'].columns)
ok(len(recap['df']) == mr.MAX_DETAILS + 15, len(recap['df']))
ok(any(b['t'] == 'p' and f"Détail des {mr.MAX_DETAILS} premiers indicateurs sur {mr.MAX_DETAILS + 15}" in b['text'] for b in blocks))
print("OK — plafond de détail, récapitulatif complet")

with m.app.test_client() as c:
    # ── Refus propres
    with c.session_transaction() as sess:
        sess['multi_surveys'] = entries[:1]
        sess['ms_auto_mode'] = True
    r = c.get('/multi-survey/report.pdf')
    ok(r.status_code == 302 and r.location.endswith('/multi-survey'), (r.status_code, r.location))
    ok('au moins deux enquêtes' in c.get('/multi-survey').get_data(as_text=True))
    with c.session_transaction() as sess:
        sess['multi_surveys'] = entries
        sess['ms_auto_mode'] = False
        sess['ms_compared_vars'] = []
    r = c.get('/multi-survey/report.docx')
    ok(r.status_code == 302, r.status_code)
    page = c.get('/multi-survey').get_data(as_text=True)
    ok("Aucune comparaison à rapporter" in page and 'id="msReport"' in page and 'Rien à rapporter' in page, 'message de refus / carte vide')
    ok(c.get('/multi-survey/report.odt').status_code == 404, 'format inconnu')
    print("OK — refus propres : moins de deux enquêtes, aucune comparaison")

    # ── Mode automatique : les trois formats
    with c.session_transaction() as sess:
        sess['ms_auto_mode'] = True
        sess['user_comments'] = {'multi_survey:Fumeur__B': COMMENT}
    page = c.get('/multi-survey').get_data(as_text=True)
    ok('6. Générer le rapport' in page and '/multi-survey/report.pdf' in page and '/multi-survey/report.docx' in page
       and '/multi-survey/report.xlsx' in page, 'boutons du rapport absents')

    q = '?title=Rapport%20CAP%20Tabac&author=%C3%89quipe%20test'
    r = c.get('/multi-survey/report.docx' + q)
    ok(r.status_code == 200 and 'wordprocessingml' in r.mimetype and r.data[:2] == b'PK', (r.status_code, r.mimetype))
    ok('rapport_multi-enquete_' in r.headers['Content-Disposition'] and r.headers['Content-Disposition'].endswith('.docx'))
    from docx import Document
    doc = Document(io.BytesIO(r.data))
    text = '\n'.join(p.text for p in doc.paragraphs)
    ok('Rapport CAP Tabac' in text and 'Synthèse' in text and 'Notes de méthode' in text, 'structure du Word')
    ok(text.count('Lecture automatique') == 4, text.count('Lecture automatique'))
    ok(f"Commentaire de l'analyste : {COMMENT}" in text, 'commentaire libre absent du Word')
    titles = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2' and p.text.startswith('4.')]
    ok([t.split('  ')[1] for t in titles] == ['Fumeur__B', 'Age__1Y', 'Region__X', 'Sexe__X'], titles)   # établis d'abord, stables ensuite
    cover = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    ok(cover['Auteur'] == 'Équipe test' and 'Enquête_2023.xlsx → Enquête_2024.xlsx → Enquête_2025.xlsx' in cover['Enquêtes comparées'], cover)
    ok(any('Unités différentes' in cell.text for t in doc.tables for row in t.rows for cell in row.cells), 'contrôle de cohérence absent')
    if charts_available:
        ok(len(doc.inline_shapes) == 4, len(doc.inline_shapes))
    print("OK — Word : couverture, pile, contrôles, synthèse, 4 comparaisons classées, lecture, commentaire, graphiques")

    from openpyxl import load_workbook
    r = c.get('/multi-survey/report.xlsx' + q)
    ok(r.status_code == 200 and 'spreadsheetml' in r.mimetype, (r.status_code, r.mimetype))
    wb = load_workbook(io.BytesIO(r.data))
    ok(wb.sheetnames == ['Résumé', 'Comparaisons', 'Méthode'], wb.sheetnames)
    cells = [str(v) for ws in wb for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    ok('Rapport CAP Tabac' in cells and any('Synthèse' in v for v in cells), 'résumé Excel')
    ok(cells.count('Lecture automatique') == 4 and any(COMMENT in v for v in cells), 'lecture / commentaire dans Excel')
    if charts_available:
        ok(len(wb['Comparaisons']._images) == 4, len(wb['Comparaisons']._images))
    print("OK — Excel : feuilles Résumé / Comparaisons / Méthode, lecture et commentaire")

    r = c.get('/multi-survey/report.pdf' + q)
    ok(r.status_code == 200 and r.mimetype == 'application/pdf' and r.data[:4] == b'%PDF', (r.status_code, r.mimetype))
    if shutil.which('pdftotext'):
        pdf = os.path.join(work, 'r.pdf')
        open(pdf, 'wb').write(r.data)
        out = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True).stdout
        flat = ' '.join(out.split())
        ok('Rapport CAP Tabac' in flat and '1. Enquêtes comparées' in flat and '2. Contrôles de cohérence' in flat, 'structure du PDF')
        ok(flat.count('Lecture automatique') == 4 and 'Hausse cohérente avec la campagne de 2024 <à confirmer> & suivi.' in flat, 'lecture / commentaire dans le PDF')
        ok(flat.index('4.1 Fumeur__B') < flat.index('4.2 Age__1Y') < flat.index('4.3 Region__X') < flat.index('4.4 Sexe__X'), 'ordre des indicateurs dans le PDF')
        ok('5. Notes de méthode' in flat and 'Benjamini-Hochberg' in flat)
    print("OK — PDF : structure, lecture et commentaire (« < » et « & » compris)")

    # ── Panier manuel : seules les variables du panier sont rapportées
    with c.session_transaction() as sess:
        sess['ms_auto_mode'] = False
        sess['ms_compared_vars'] = ['Fumeur__B', 'Age__1Y']
    r = c.get('/multi-survey/report.docx')
    ok(r.status_code == 200, r.status_code)
    text = '\n'.join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
    ok(text.count('Lecture automatique') == 2 and 'Region__X' not in text and 'Rapport' in text, text.count('Lecture automatique'))
    cover = {row.cells[0].text: row.cells[1].text for row in Document(io.BytesIO(r.data)).tables[0].rows}
    ok(cover['Indicateurs comparés'] == '2 dont 2 avec un écart établi', cover)
    print("OK — panier manuel : uniquement les variables choisies")

print()
print(f"TOUS LES TESTS DU RAPPORT MULTI-ENQUÊTE SONT PASSÉS ({checks} vérifications)")
